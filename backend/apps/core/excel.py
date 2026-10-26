"""Exportación a Excel (.xlsx) con encabezado institucional.

Uso:
    return respuesta_excel(
        "interesados.xlsx",
        titulo="Interesados — Convocatoria 2026-2027",
        columnas=[("Nombre", 30), ("CI", 14)],
        filas=[["María", "98042512345"], ...],
    )
"""

from io import BytesIO

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FILA_ENCABEZADO_TABLA = 5


def _institucion():
    from apps.configuracion.models import IdentidadVisual

    identidad = IdentidadVisual.actual()
    return identidad.nombre_institucion, identidad.color_primario.lstrip("#")


def libro_excel(titulo, columnas, filas):
    """Libro con: institución y sistema, título, fecha de generación y la tabla con filtros."""
    institucion, color = _institucion()
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Datos"

    hoja["A1"] = institucion
    hoja["A1"].font = Font(bold=True, size=14, color=color)
    hoja["A2"] = "SIGECUE — Sistema de Gestión de Carreras por Encuentro"
    hoja["A3"] = titulo
    hoja["A3"].font = Font(bold=True, size=12)
    hoja["A4"] = f"Generado el {timezone.localtime():%d/%m/%Y %H:%M}"
    hoja["A4"].font = Font(italic=True, color="6B7280")

    relleno = PatternFill("solid", fgColor=color)
    for indice, (encabezado, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(row=FILA_ENCABEZADO_TABLA, column=indice, value=encabezado)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    for fila in filas:
        hoja.append(list(fila))

    ultima_columna = get_column_letter(len(columnas))
    hoja.auto_filter.ref = f"A{FILA_ENCABEZADO_TABLA}:{ultima_columna}{max(hoja.max_row, FILA_ENCABEZADO_TABLA)}"
    hoja.freeze_panes = f"A{FILA_ENCABEZADO_TABLA + 1}"
    return libro


def respuesta_excel(nombre_archivo, titulo, columnas, filas):
    contenido = BytesIO()
    libro_excel(titulo, columnas, filas).save(contenido)
    respuesta = HttpResponse(contenido.getvalue(), content_type=TIPO_XLSX)
    respuesta["Content-Disposition"] = f'attachment; filename="{nombre_archivo}"'
    return respuesta
