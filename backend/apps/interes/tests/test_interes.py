from datetime import date, timedelta
from io import BytesIO

import pytest
from django.core import mail
from django.db import IntegrityError, connection
from django.utils import timezone
from openpyxl import load_workbook

from apps.auditoria.models import LogAuditoria
from apps.interes import services
from apps.interes.models import ConvocatoriaInteres, ExpresionInteres
from apps.interes.tasks import cerrar_convocatorias_vencidas
from apps.nomencladores.models import Carrera
from apps.oferta.models import OfertaCarrera

pytestmark = pytest.mark.django_db

ADMIN = "/api/v1/admin/interes"
HOY = timezone.localdate()


@pytest.fixture
def jefe_ci(crear_usuario):
    return crear_usuario("jefe", rol="jefe_ci")


@pytest.fixture
def admin_facultad(crear_usuario, facultad):
    return crear_usuario("admin_fac", rol="admin_facultad", facultad=facultad)


@pytest.fixture
def informatica(db):
    return Carrera.objects.create(nombre="Ingeniería Informática", codigo="II", duracion_anios=5, descripcion="Software")


@pytest.fixture
def derecho(db):
    return Carrera.objects.create(nombre="Derecho", codigo="DER", duracion_anios=5)


@pytest.fixture
def convocatoria(db):
    return ConvocatoriaInteres.objects.create(
        anio_academico="2026-2027", fecha_inicio=HOY - timedelta(days=1), fecha_cierre=HOY + timedelta(days=30)
    )


@pytest.fixture
def oferta(convocatoria, informatica, facultad):
    return OfertaCarrera.objects.create(carrera=informatica, facultad=facultad, anio_academico="2026-2027", asignada_por_ci=True)


@pytest.fixture
def abierta(convocatoria, oferta):
    services.abrir(convocatoria)
    return convocatoria


def interes(**extra):
    return {
        "nombre": "María González Pérez", "ci": "98042512345", "celular": "55123456",
        "email": "maria@nauta.cu", "acepta_privacidad": True,
    } | extra


class TestConvocatoria:
    def test_la_crea_la_comision(self, api_como, jefe_ci):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/convocatorias/",
            {"anio_academico": "2026-2027", "fecha_inicio": "2026-11-01", "fecha_cierre": "2026-12-15"},
            format="json",
        )

        assert response.status_code == 201
        assert response.json()["estado"] == "borrador"

    def test_rol_local_no_crea_convocatorias(self, api_como, admin_facultad):
        response = api_como(admin_facultad).post(f"{ADMIN}/convocatorias/", {"anio_academico": "2026-2027"}, format="json")

        assert response.status_code == 403
        assert "Comisión de Ingreso" in response.json()["error"]["mensaje"]

    @pytest.mark.parametrize("anio", ["2026", "2026-2028", "26-27"])
    def test_formato_del_anio(self, api_como, jefe_ci, anio):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/convocatorias/", {"anio_academico": anio, "fecha_inicio": "2026-11-01", "fecha_cierre": "2026-12-01"}, format="json"
        )

        assert "anio_academico" in response.json()["error"]["detalles"]

    def test_cierre_no_anterior_al_inicio(self, api_como, jefe_ci):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/convocatorias/",
            {"anio_academico": "2026-2027", "fecha_inicio": "2026-12-01", "fecha_cierre": "2026-11-01"},
            format="json",
        )

        assert "fecha_cierre" in response.json()["error"]["detalles"]

    def test_una_por_anio(self, api_como, jefe_ci, convocatoria):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/convocatorias/", {"anio_academico": "2026-2027", "fecha_inicio": "2026-11-01", "fecha_cierre": "2026-12-01"}, format="json"
        )

        assert "anio_academico" in response.json()["error"]["detalles"]

    def test_abrir_exige_oferta(self, api_como, jefe_ci, convocatoria):
        response = api_como(jefe_ci).post(f"{ADMIN}/convocatorias/{convocatoria.pk}/abrir/")

        assert "al menos una carrera" in response.json()["error"]["mensaje"]

    def test_abrir_y_cerrar(self, api_como, jefe_ci, convocatoria, oferta):
        client = api_como(jefe_ci)

        assert client.post(f"{ADMIN}/convocatorias/{convocatoria.pk}/abrir/").json()["estado"] == "abierta"
        assert client.post(f"{ADMIN}/convocatorias/{convocatoria.pk}/cerrar/").json()["estado"] == "cerrada"

    def test_solo_una_abierta(self, abierta, informatica, facultad):
        otra = ConvocatoriaInteres.objects.create(anio_academico="2027-2028", fecha_inicio=HOY, fecha_cierre=HOY + timedelta(days=5))
        OfertaCarrera.objects.create(carrera=informatica, facultad=facultad, anio_academico="2027-2028")

        with pytest.raises(Exception, match="otra convocatoria"):
            services.abrir(otra)

    def test_la_base_de_datos_impide_dos_abiertas(self, abierta):
        with pytest.raises(IntegrityError):
            ConvocatoriaInteres.objects.create(
                anio_academico="2027-2028", fecha_inicio=HOY, fecha_cierre=HOY, estado=ConvocatoriaInteres.Estado.ABIERTA
            )

    def test_cierre_automatico_por_fecha(self, abierta):
        ConvocatoriaInteres.objects.filter(pk=abierta.pk).update(fecha_cierre=HOY - timedelta(days=1))

        assert cerrar_convocatorias_vencidas() == 1
        abierta.refresh_from_db()
        assert abierta.estado == "cerrada"

    def test_cerrada_no_se_modifica(self, api_como, jefe_ci, abierta):
        services.cerrar(abierta)

        response = api_como(jefe_ci).patch(f"{ADMIN}/convocatorias/{abierta.pk}/", {"mensaje_publico": "x"}, format="json")

        assert response.status_code == 400

    def test_solo_se_elimina_en_borrador(self, api_como, jefe_ci, abierta, convocatoria):
        client = api_como(jefe_ci)

        assert client.delete(f"{ADMIN}/convocatorias/{abierta.pk}/").status_code == 400
        services.cerrar(abierta)
        borrador = ConvocatoriaInteres.objects.create(anio_academico="2027-2028", fecha_inicio=HOY, fecha_cierre=HOY)
        assert client.delete(f"{ADMIN}/convocatorias/{borrador.pk}/").status_code == 204


class TestOferta:
    def test_la_comision_asigna_a_cualquier_entidad(self, api_como, jefe_ci, convocatoria, informatica, cum):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/ofertas/", {"convocatoria": convocatoria.pk, "carrera": informatica.pk, "cum_fum": cum.pk}, format="json"
        )

        assert response.status_code == 201
        assert response.json()["asignada_por_ci"] is True
        assert response.json()["entidad"]["nombre"] == "CUM Calixto García"

    def test_la_entidad_agrega_a_la_suya_sin_indicarla(self, api_como, admin_facultad, convocatoria, derecho, facultad):
        response = api_como(admin_facultad).post(f"{ADMIN}/ofertas/", {"convocatoria": convocatoria.pk, "carrera": derecho.pk}, format="json")

        assert response.status_code == 201
        oferta = OfertaCarrera.objects.get(carrera=derecho)
        assert oferta.facultad == facultad and not oferta.asignada_por_ci

    def test_la_entidad_no_agrega_a_otra(self, api_como, admin_facultad, convocatoria, derecho, otra_facultad):
        response = api_como(admin_facultad).post(
            f"{ADMIN}/ofertas/", {"convocatoria": convocatoria.pk, "carrera": derecho.pk, "facultad": otra_facultad.pk}, format="json"
        )

        assert "propia entidad" in response.json()["error"]["mensaje"]

    def test_no_se_repite(self, api_como, jefe_ci, convocatoria, oferta, informatica, facultad):
        response = api_como(jefe_ci).post(
            f"{ADMIN}/ofertas/", {"convocatoria": convocatoria.pk, "carrera": informatica.pk, "facultad": facultad.pk}, format="json"
        )

        assert "ya está en la oferta" in response.json()["error"]["mensaje"]

    def test_la_misma_carrera_en_facultad_y_cum(self, jefe_ci, convocatoria, oferta, informatica, cum):
        services.agregar_oferta(convocatoria, informatica, None, cum, jefe_ci)

        assert OfertaCarrera.objects.filter(carrera=informatica).count() == 2

    def test_la_entidad_no_quita_lo_asignado_por_la_ci(self, api_como, admin_facultad, oferta):
        response = api_como(admin_facultad).delete(f"{ADMIN}/ofertas/{oferta.pk}/")

        assert "Comisión de Ingreso" in response.json()["error"]["mensaje"]

    def test_la_entidad_quita_lo_suyo(self, api_como, admin_facultad, convocatoria, derecho):
        propia = services.agregar_oferta(convocatoria, derecho, None, None, admin_facultad)

        assert api_como(admin_facultad).delete(f"{ADMIN}/ofertas/{propia.pk}/").status_code == 204

    def test_no_se_quita_con_interesados(self, api_como, jefe_ci, abierta, oferta):
        services.registrar_interes(oferta, "Ana", "98042512345", "55123456", "a@b.cu")

        assert "expresiones de interés" in api_como(jefe_ci).delete(f"{ADMIN}/ofertas/{oferta.pk}/").json()["error"]["mensaje"]

    def test_cerrada_no_se_modifica(self, api_como, jefe_ci, abierta, derecho, facultad):
        services.cerrar(abierta)

        response = api_como(jefe_ci).post(
            f"{ADMIN}/ofertas/", {"convocatoria": abierta.pk, "carrera": derecho.pk, "facultad": facultad.pk}, format="json"
        )

        assert "cerrada" in response.json()["error"]["mensaje"]

    def test_rol_local_ve_la_oferta_de_su_entidad(self, api_como, admin_facultad, jefe_ci, convocatoria, oferta, derecho, cum):
        services.agregar_oferta(convocatoria, derecho, None, cum, jefe_ci)

        locales = api_como(admin_facultad).get(f"{ADMIN}/ofertas/", {"convocatoria": convocatoria.pk}).json()["results"]

        assert [o["carrera_nombre"] for o in locales] == ["Ingeniería Informática"]


class TestRegistroPublico:
    url = "/api/v1/interes/"

    def test_registra_sin_cuenta_y_envia_confirmacion(self, api, abierta, oferta, django_capture_on_commit_callbacks):
        with django_capture_on_commit_callbacks(execute=True):
            response = api.post(self.url, interes(oferta=oferta.pk), format="json")

        assert response.status_code == 201
        expresion = ExpresionInteres.objects.get()
        assert (expresion.nombre, expresion.ci) == ("María González Pérez", "98042512345")
        assert mail.outbox[0].to == ["maria@nauta.cu"]
        assert "Ingeniería Informática" in mail.outbox[0].subject

    def test_datos_personales_cifrados(self, api, abierta, oferta):
        api.post(self.url, interes(oferta=oferta.pk), format="json")

        with connection.cursor() as cursor:
            cursor.execute("SELECT ci, celular, email FROM interes_expresioninteres")
            fila = " ".join(cursor.fetchone())
        assert "98042512345" not in fila and "55123456" not in fila and "maria@nauta.cu" not in fila

    def test_mismo_ci_no_repite_la_misma_oferta(self, api, abierta, oferta):
        api.post(self.url, interes(oferta=oferta.pk), format="json")

        response = api.post(self.url, interes(oferta=oferta.pk, email="otro@nauta.cu"), format="json")

        assert "Ya registraste" in response.json()["error"]["mensaje"]

    def test_mismo_ci_en_otra_oferta(self, api, abierta, oferta, informatica, cum, jefe_ci):
        otra = services.agregar_oferta(abierta, informatica, None, cum, jefe_ci)
        api.post(self.url, interes(oferta=oferta.pk), format="json")

        assert api.post(self.url, interes(oferta=otra.pk), format="json").status_code == 201

    @pytest.mark.parametrize("campo,valor", [("ci", "123"), ("celular", "44123456"), ("email", "no-es-email"), ("nombre", "A")])
    def test_validaciones(self, api, abierta, oferta, campo, valor):
        response = api.post(self.url, interes(oferta=oferta.pk, **{campo: valor}), format="json")

        assert campo in response.json()["error"]["detalles"]

    def test_exige_aceptar_la_privacidad(self, api, abierta, oferta):
        response = api.post(self.url, interes(oferta=oferta.pk, acepta_privacidad=False), format="json")

        assert "acepta_privacidad" in response.json()["error"]["detalles"]

    def test_sin_convocatoria_abierta_no_acepta(self, api, convocatoria, oferta):
        response = api.post(self.url, interes(oferta=oferta.pk), format="json")

        assert "no está abierta" in response.json()["error"]["mensaje"]

    def test_abierta_pero_aun_no_empieza(self, api, abierta, oferta):
        ConvocatoriaInteres.objects.filter(pk=abierta.pk).update(fecha_inicio=HOY + timedelta(days=2), fecha_cierre=HOY + timedelta(days=9))

        assert api.post(self.url, interes(oferta=oferta.pk), format="json").status_code == 400

    def test_oferta_de_otro_anio(self, api, abierta, informatica, facultad):
        vieja = OfertaCarrera.objects.create(carrera=informatica, facultad=facultad, anio_academico="2025-2026")

        assert "no está abierta" in api.post(self.url, interes(oferta=vieja.pk), format="json").json()["error"]["mensaje"]

    def test_limite_por_ip(self, api, abierta, oferta, settings):
        settings.REST_FRAMEWORK = settings.REST_FRAMEWORK | {
            "DEFAULT_THROTTLE_RATES": {**settings.REST_FRAMEWORK["DEFAULT_THROTTLE_RATES"], "interes": "2/hour"}
        }
        from rest_framework.throttling import ScopedRateThrottle

        ScopedRateThrottle.THROTTLE_RATES = settings.REST_FRAMEWORK["DEFAULT_THROTTLE_RATES"]
        try:
            codigos = [api.post(self.url, interes(oferta=oferta.pk, ci=f"9804251234{i}"), format="json").status_code for i in range(3)]
        finally:
            from rest_framework.settings import api_settings

            ScopedRateThrottle.THROTTLE_RATES = api_settings.DEFAULT_THROTTLE_RATES

        assert codigos == [201, 201, 429]


class TestPublico:
    def test_convocatoria_activa(self, api, abierta):
        assert api.get("/api/v1/interes/convocatoria/").json()["convocatoria"]["anio_academico"] == "2026-2027"

    def test_sin_convocatoria_activa(self, api, convocatoria):
        assert api.get("/api/v1/interes/convocatoria/").json() == {"convocatoria": None}

    def test_oferta_publica_y_filtros(self, api, abierta, oferta, jefe_ci, derecho, cum):
        services.agregar_oferta(abierta, derecho, None, cum, jefe_ci)

        todas = api.get("/api/v1/oferta/").json()
        assert todas["count"] == 2
        assert set(todas["results"][0]) == {"id", "carrera", "entidad", "anio_academico"}
        assert api.get("/api/v1/oferta/", {"cum_fum": cum.pk}).json()["count"] == 1
        assert api.get("/api/v1/oferta/", {"municipio": cum.municipio_id}).json()["count"] == 1
        assert api.get("/api/v1/oferta/", {"buscar": "infor"}).json()["count"] == 1

    def test_sin_convocatoria_la_oferta_esta_vacia(self, api, convocatoria, oferta):
        assert api.get("/api/v1/oferta/").json()["count"] == 0

    def test_catalogo_de_carreras_activas(self, api, informatica, derecho):
        Carrera.objects.filter(pk=derecho.pk).update(activo=False)

        assert [c["nombre"] for c in api.get("/api/v1/carreras/").json()["results"]] == ["Ingeniería Informática"]

    def test_detalle_con_donde_se_oferta(self, api, abierta, oferta, informatica):
        datos = api.get(f"/api/v1/carreras/{informatica.pk}/").json()

        assert datos["descripcion"] == "Software"
        assert datos["ofertas"][0]["entidad"]["nombre"] == "Facultad de Informática"

    def test_filtros_publicos(self, api, facultad, cum):
        datos = api.get("/api/v1/publico/filtros/").json()

        assert {f["nombre"] for f in datos["facultades"]} == {"Facultad de Informática"}
        assert len(datos["municipios"]) == 14

    def test_no_exige_ni_rompe_con_token_invalido(self, api, abierta):
        api.credentials(HTTP_AUTHORIZATION="Bearer token-caducado")

        assert api.get("/api/v1/oferta/").status_code == 200


class TestInteresados:
    @pytest.fixture
    def expresiones(self, abierta, oferta, jefe_ci, derecho, cum):
        otra = services.agregar_oferta(abierta, derecho, None, cum, jefe_ci)
        services.registrar_interes(oferta, "María", "98042512345", "55123456", "maria@nauta.cu")
        services.registrar_interes(otra, "Luis", "01081923456", "55999999", "luis@nauta.cu")

    def test_rol_local_ve_solo_los_de_su_entidad(self, api_como, admin_facultad, expresiones):
        datos = api_como(admin_facultad).get(f"{ADMIN}/expresiones/").json()

        assert [e["nombre"] for e in datos["results"]] == ["María"]
        assert datos["results"][0]["ci"] == "98042512345"

    def test_rol_global_ve_todos_y_filtra(self, api_como, jefe_ci, expresiones, cum):
        client = api_como(jefe_ci)

        assert client.get(f"{ADMIN}/expresiones/").json()["count"] == 2
        assert client.get(f"{ADMIN}/expresiones/", {"cum_fum": cum.pk}).json()["count"] == 1
        assert client.get(f"{ADMIN}/expresiones/", {"buscar": "lui"}).json()["count"] == 1

    def test_exporta_excel_y_queda_auditado(self, api_como, jefe_ci, expresiones, cum):
        response = api_como(jefe_ci).get(f"{ADMIN}/expresiones/exportar/", {"cum_fum": cum.pk})

        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")
        hoja = load_workbook(BytesIO(response.content)).active
        assert hoja["A1"].value == "Universidad de Holguín"
        assert [c.value for c in hoja[5]][:4] == ["Nombre", "CI", "Celular", "Email"]
        assert [c.value for c in hoja[6]][:2] == ["Luis", "01081923456"]
        assert hoja.max_row == 6
        log = LogAuditoria.objects.get(accion="exportar interesados")
        assert log.detalle == {"filtros": {"cum_fum": str(cum.pk)}, "registros": 1}

    def test_exportar_requiere_permiso(self, api_como, crear_usuario, facultad, expresiones):
        jefe_area = crear_usuario("ja", rol="jefe_area_facultad", facultad=facultad)

        assert api_como(jefe_area).get(f"{ADMIN}/expresiones/exportar/").status_code == 403


def test_str(oferta, convocatoria):
    assert str(convocatoria) == "Convocatoria de interés 2026-2027"
    assert str(oferta) == "Ingeniería Informática — Facultad de Informática (2026-2027)"
    assert date.fromisoformat("2026-01-01")
