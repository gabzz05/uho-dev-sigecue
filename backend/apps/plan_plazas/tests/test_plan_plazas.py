from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.core.pdf import pdf_disponible
from apps.nomencladores.models import Area, Carrera
from apps.notificaciones.models import Notificacion
from apps.oferta.models import OfertaCarrera
from apps.plan_plazas import services
from apps.plan_plazas.models import HistorialPlan, LineaPlan, PlanEntidad, PlanPlazas

pytestmark = pytest.mark.django_db

URL = "/api/v1/admin/plan-plazas/"
LINEAS = "/api/v1/admin/plan-plazas/lineas/"
ANIO = "2026-2027"


@pytest.fixture
def jefe_ci(crear_usuario):
    return crear_usuario("jefe", rol="jefe_ci", first_name="Juan")


@pytest.fixture
def vicerrector(crear_usuario):
    return crear_usuario("vice", rol="vicerrector", first_name="Rosa", last_name="Díaz")


@pytest.fixture
def admin_fac(crear_usuario, facultad):
    return crear_usuario("admin_fac", rol="admin_facultad", facultad=facultad)


@pytest.fixture
def vicedecano(crear_usuario, facultad):
    return crear_usuario("vicedecano", rol="vicedecano", facultad=facultad)


@pytest.fixture
def jefe_area(crear_usuario, facultad):
    area = Area.objects.create(nombre="Software", facultad=facultad)
    return crear_usuario("jefe_area", rol="jefe_area_facultad", facultad=facultad, area=area)


@pytest.fixture
def admin_cum(crear_usuario, cum):
    return crear_usuario("admin_cum", rol="admin_cum", cum_fum=cum)


@pytest.fixture
def carreras(db):
    return {
        codigo: Carrera.objects.create(nombre=nombre, codigo=codigo, duracion_anios=5)
        for codigo, nombre in [("II", "Ingeniería Informática"), ("DER", "Derecho"), ("CF", "Contabilidad y Finanzas")]
    }


@pytest.fixture
def plan(jefe_ci, carreras, facultad, otra_facultad, cum):
    # Todas las entidades existen antes de enviar: el envío llega a las activas en ese momento.
    plan = services.crear_plan(ANIO, jefe_ci)
    services.agregar_linea_ci(plan, carreras["II"], facultad, None, 80, jefe_ci)
    services.agregar_linea_ci(plan, carreras["DER"], facultad, None, 60, jefe_ci)
    services.agregar_linea_ci(plan, carreras["DER"], None, cum, 40, jefe_ci)
    return plan


@pytest.fixture
def enviado(plan, jefe_ci, admin_fac, vicedecano, admin_cum):
    return services.enviar(plan, jefe_ci)


def linea(plan, codigo, entidad):
    filtro = {"facultad": entidad} if hasattr(entidad, "decano") else {"cum_fum": entidad}
    return plan.lineas.get(carrera__codigo=codigo, **filtro)


def errores(response):
    return response.json()["error"]


class TestBorrador:
    def test_la_ci_crea_el_plan(self, api_como, jefe_ci):
        response = api_como(jefe_ci).post(URL, {"anio_academico": ANIO}, format="json")

        assert response.status_code == 201
        assert response.json()["estado"] == "borrador"
        assert "enviar" in response.json()["acciones"]

    def test_un_plan_base_por_anio(self, api_como, jefe_ci, plan):
        response = api_como(jefe_ci).post(URL, {"anio_academico": ANIO}, format="json")

        assert "versión complementaria" in errores(response)["mensaje"]

    def test_un_rol_local_no_crea_planes(self, api_como, admin_fac):
        # No tiene plan_plazas.crear; aunque lo tuviera, el service exige ser de la CI.
        assert api_como(admin_fac).post(URL, {"anio_academico": ANIO}, format="json").status_code == 403

    def test_cargar_la_oferta_del_anio(self, api_como, jefe_ci, plan, carreras, facultad):
        OfertaCarrera.objects.create(carrera=carreras["CF"], facultad=facultad, anio_academico=ANIO)
        OfertaCarrera.objects.create(carrera=carreras["II"], facultad=facultad, anio_academico=ANIO)  # Ya está en el plan.

        api_como(jefe_ci).post(f"{URL}{plan.pk}/cargar-oferta/")

        assert plan.lineas.count() == 4
        assert linea(plan, "CF", facultad).plazas_propuestas == 0

    def test_editar_y_quitar_lineas(self, api_como, jefe_ci, plan, facultad):
        client = api_como(jefe_ci)
        informatica = linea(plan, "II", facultad)

        assert client.patch(f"{LINEAS}{informatica.pk}/", {"plazas_propuestas": 90}, format="json").status_code == 200
        assert client.delete(f"{LINEAS}{informatica.pk}/").status_code == 204
        assert plan.lineas.count() == 2

    def test_linea_repetida(self, api_como, jefe_ci, plan, carreras, facultad):
        response = api_como(jefe_ci).post(
            LINEAS, {"plan": plan.pk, "carrera": carreras["II"].pk, "facultad": facultad.pk, "plazas": 5}, format="json"
        )

        assert "ya está en el plan" in errores(response)["mensaje"]

    def test_enviar_exige_lineas(self, jefe_ci):
        vacio = services.crear_plan("2027-2028", jefe_ci)

        with pytest.raises(Exception, match="al menos una carrera"):
            services.enviar(vacio, jefe_ci)


class TestEnvio:
    def test_llega_a_todas_las_entidades_y_notifica_a_quienes_responden(self, enviado, otra_facultad, jefe_area, admin_fac, vicedecano):
        assert enviado.estado == "enviado"
        assert enviado.entidades.count() == 3  # Dos facultades y un CUM, con o sin líneas.
        notificados = set(Notificacion.objects.filter(tipo="plan_enviado").values_list("usuario__username", flat=True))
        assert {"admin_fac", "vicedecano", "admin_cum"} <= notificados
        assert "jefe_area" not in notificados

    def test_el_borrador_queda_cerrado(self, enviado, jefe_ci, facultad):
        with pytest.raises(Exception, match="Borrador"):
            services.actualizar_linea_ci(linea(enviado, "II", facultad), 10, jefe_ci)


class TestRevisionDeLaEntidad:
    def test_modificar_exige_cantidad_y_comentario(self, api_como, admin_fac, enviado, facultad):
        url = f"{LINEAS}{linea(enviado, 'II', facultad).pk}/responder/"
        client = api_como(admin_fac)

        assert "cantidad" in errores(client.post(url, {"accion": "modificar"}, format="json"))["mensaje"]
        assert "comentario" in errores(client.post(url, {"accion": "modificar", "plazas": 100}, format="json"))["mensaje"]
        datos = client.post(url, {"accion": "modificar", "plazas": 100, "comentario": "Hay demanda"}, format="json").json()

        assert datos["propuesta_entidad"] == 100
        assert PlanEntidad.objects.get(plan=enviado, facultad=facultad).estado == "en_revision"

    def test_quitar_exige_comentario(self, admin_fac, enviado, facultad):
        with pytest.raises(Exception, match="comentario"):
            services.responder_linea(linea(enviado, "DER", facultad), "quitar", None, "", admin_fac)

    def test_no_toca_lineas_de_otra_entidad(self, api_como, admin_fac, enviado, cum):
        response = api_como(admin_fac).post(f"{LINEAS}{linea(enviado, 'DER', cum).pk}/responder/", {"accion": "aceptar"}, format="json")

        assert response.status_code == 404

    def test_solicitar_y_retirar_una_carrera(self, api_como, admin_fac, enviado, carreras):
        client = api_como(admin_fac)
        datos = {"plan": enviado.pk, "carrera": carreras["CF"].pk, "plazas": 30}

        assert "comentario" in errores(client.post(LINEAS, datos, format="json"))["mensaje"]
        nueva = client.post(LINEAS, datos | {"comentario": "Nueva demanda"}, format="json").json()
        assert (nueva["origen"], nueva["accion_entidad"], nueva["propuesta_entidad"]) == ("entidad", "agregar", 30)
        assert client.delete(f"{LINEAS}{nueva['id']}/").status_code == 204

    def test_el_jefe_de_area_comenta_pero_no_responde(self, api_como, jefe_area, enviado, facultad):
        client = api_como(jefe_area)
        informatica = linea(enviado, "II", facultad)

        respuesta = client.post(f"{LINEAS}{informatica.pk}/responder/", {"accion": "aceptar"}, format="json")
        comentario = client.post(f"{URL}{enviado.pk}/comentarios/", {"texto": "Sugiero 90", "linea": informatica.pk}, format="json")

        assert "Admin, el Vicedecano o el Director" in errores(respuesta)["mensaje"]
        assert comentario.status_code == 201
        assert "Jefe de Área" in comentario.json()["usuario_nombre"]

    def test_los_comentarios_se_ven_por_ambito(self, api_como, jefe_area, admin_cum, jefe_ci, enviado):
        services.comentar(enviado, jefe_area, "Desde la facultad")

        assert len(api_como(admin_cum).get(f"{URL}{enviado.pk}/comentarios/").json()) == 0
        assert len(api_como(jefe_ci).get(f"{URL}{enviado.pk}/comentarios/").json()) == 1

    def test_enviar_la_propuesta_la_bloquea(self, api_como, vicedecano, enviado, facultad):
        client = api_como(vicedecano)
        client.post(f"{URL}{enviado.pk}/enviar-propuesta/", {"texto": "Conforme"}, format="json")

        response = client.post(f"{LINEAS}{linea(enviado, 'II', facultad).pk}/responder/", {"accion": "aceptar"}, format="json")

        assert "ya envió su propuesta" in errores(response)["mensaje"]
        entidad = PlanEntidad.objects.get(plan=enviado, facultad=facultad)
        assert (entidad.estado, entidad.enviada_por, entidad.comentario) == ("propuesta_enviada", vicedecano, "Conforme")

    def test_los_locales_ven_solo_sus_lineas(self, api_como, admin_fac, enviado):
        datos = api_como(admin_fac).get(LINEAS, {"plan": enviado.pk}).json()

        assert {l["entidad"]["nombre"] for l in datos} == {"Facultad de Informática"}

    def test_acciones_segun_el_rol(self, api_como, admin_fac, jefe_area, jefe_ci, enviado):
        def acciones(usuario):
            return api_como(usuario).get(f"{URL}{enviado.pk}/").json()["acciones"]

        assert {"responder", "enviar_propuesta"} <= set(acciones(admin_fac))
        assert "responder" not in acciones(jefe_area) and "comentar" in acciones(jefe_area)
        assert "cerrar_revision" in acciones(jefe_ci)


@pytest.fixture
def con_propuestas(enviado, admin_fac, admin_cum, carreras, facultad, cum, otra_facultad, crear_usuario):
    """Facultad: sube Informática a 100, pide quitar Derecho y agregar Contabilidad. El CUM no responde."""
    services.responder_linea(linea(enviado, "II", facultad), "modificar", 100, "Hay demanda", admin_fac)
    services.responder_linea(linea(enviado, "DER", facultad), "quitar", None, "Sin profesores", admin_fac)
    services.solicitar_carrera(enviado, carreras["CF"], 30, "Nueva demanda", admin_fac)
    services.enviar_propuesta(enviado, admin_fac)
    services.responder_linea(linea(enviado, "DER", cum), "modificar", 10, "Borrador sin enviar", admin_cum)
    admin_otra = crear_usuario("admin_otra", rol="admin_facultad", facultad=otra_facultad)
    services.enviar_propuesta(enviado, admin_otra)
    return enviado


class TestRevisionDeLaCI:
    def test_esperar_o_forzar(self, api_como, jefe_ci, con_propuestas, cum):
        client = api_como(jefe_ci)

        sin_forzar = client.post(f"{URL}{con_propuestas.pk}/cerrar-revision/", {}, format="json")
        assert "CUM Calixto García" in errores(sin_forzar)["mensaje"]

        forzado = client.post(f"{URL}{con_propuestas.pk}/cerrar-revision/", {"forzar": True}, format="json").json()
        assert forzado["estado"] == "revision_ci"
        # El CUM no envió: conserva la propuesta inicial aunque hubiera empezado a editar.
        der_cum = linea(con_propuestas, "DER", cum)
        assert (der_cum.accion_entidad, der_cum.plazas_entidad) == ("aceptar", None)

    def test_decidir_exige_justificacion(self, jefe_ci, con_propuestas, facultad):
        services.cerrar_revision(con_propuestas, jefe_ci, forzar=True)
        informatica = linea(con_propuestas, "II", facultad)

        with pytest.raises(Exception, match="cantidad"):
            services.decidir(informatica, "otro", None, "x", jefe_ci)
        with pytest.raises(Exception, match="comentario"):
            services.decidir(informatica, "mantener", None, "", jefe_ci)

    def test_version_final_calcula_las_plazas(self, api_como, jefe_ci, vicerrector, con_propuestas, facultad, cum):
        services.cerrar_revision(con_propuestas, jefe_ci, forzar=True)
        services.decidir(linea(con_propuestas, "II", facultad), "otro", 90, "Capacidad máxima", jefe_ci)
        # Derecho (Facultad) pedido quitar: se acepta (por defecto). Contabilidad pedida agregar: se acepta.

        datos = api_como(jefe_ci).post(f"{URL}{con_propuestas.pk}/generar-version-final/").json()

        assert (datos["estado"], datos["version"]) == ("en_aprobacion", 1)
        aprobadas = {(l.carrera.codigo, str(l.entidad)): l.plazas_aprobadas for l in con_propuestas.lineas.all()}
        assert aprobadas == {
            ("II", "Facultad de Informática"): 90,
            ("DER", "Facultad de Informática"): None,
            ("CF", "Facultad de Informática"): 30,
            ("DER", "CUM Calixto García"): 40,
        }
        assert Notificacion.objects.filter(tipo="plan_para_aprobacion", usuario=vicerrector).exists()

    def test_mantener_una_carrera_solicitada_es_no_incluirla(self, jefe_ci, con_propuestas, facultad):
        services.cerrar_revision(con_propuestas, jefe_ci, forzar=True)
        contabilidad = linea(con_propuestas, "CF", facultad)

        services.decidir(contabilidad, "mantener", None, "No hay presupuesto", jefe_ci)

        contabilidad.refresh_from_db()
        assert contabilidad.resultado() is None


@pytest.fixture
def en_aprobacion(con_propuestas, jefe_ci):
    services.cerrar_revision(con_propuestas, jefe_ci, forzar=True)
    return services.generar_version_final(con_propuestas, jefe_ci)


class TestAprobacion:
    def test_solo_el_vicerrector_aprueba(self, api_como, jefe_ci, en_aprobacion):
        response = api_como(jefe_ci).post(f"{URL}{en_aprobacion.pk}/aprobar/")

        assert "Vicerrector" in errores(response)["mensaje"]

    def test_aprobar_publica_y_fija_la_oferta(self, api_como, vicerrector, en_aprobacion, carreras, facultad, cum, admin_fac):
        informatica_oferta = OfertaCarrera.objects.create(carrera=carreras["II"], facultad=facultad, anio_academico=ANIO)
        derecho_oferta = OfertaCarrera.objects.create(carrera=carreras["DER"], facultad=facultad, anio_academico=ANIO, plazas_plan=60)

        datos = api_como(vicerrector).post(f"{URL}{en_aprobacion.pk}/aprobar/").json()

        assert datos["estado"] == "publicado"
        assert datos["aprobado_por_nombre"] == "Rosa Díaz"
        informatica_oferta.refresh_from_db()
        derecho_oferta.refresh_from_db()
        assert informatica_oferta.plazas_plan == 100
        assert derecho_oferta.plazas_plan is None  # Salió del plan; la oferta se conserva.
        nueva = OfertaCarrera.objects.get(carrera=carreras["CF"], facultad=facultad)
        assert nueva.plazas_plan == 30 and nueva.asignada_por_ci
        assert OfertaCarrera.objects.get(carrera=carreras["DER"], cum_fum=cum).plazas_plan == 40
        assert Notificacion.objects.filter(tipo="plan_publicado", usuario=admin_fac).exists()

    def test_devolver_exige_observaciones_y_vuelve_a_la_ci(self, api_como, vicerrector, jefe_ci, en_aprobacion):
        client = api_como(vicerrector)

        assert "observaciones" in errores(client.post(f"{URL}{en_aprobacion.pk}/devolver/", {}, format="json"))["mensaje"]
        datos = client.post(f"{URL}{en_aprobacion.pk}/devolver/", {"texto": "Revisar Derecho"}, format="json").json()

        assert (datos["estado"], datos["observaciones_vicerrector"]) == ("revision_ci", "Revisar Derecho")
        assert Notificacion.objects.filter(tipo="plan_devuelto", usuario=jefe_ci).exists()
        assert services.generar_version_final(en_aprobacion, jefe_ci).version == 2

    def test_publicado_no_se_modifica(self, api_como, jefe_ci, vicerrector, en_aprobacion, facultad):
        services.aprobar(en_aprobacion, vicerrector)

        assert api_como(jefe_ci).patch(f"{URL}{en_aprobacion.pk}/", {"notas": "x"}, format="json").status_code == 400
        with pytest.raises(Exception):
            services.decidir(linea(en_aprobacion, "II", facultad), "aceptar", None, "", jefe_ci)


class TestHistorial:
    def test_registra_quien_que_antes_despues_y_version(self, api_como, jefe_ci, en_aprobacion):
        entradas = api_como(jefe_ci).get(f"{URL}{en_aprobacion.pk}/historial/").json()

        modificacion = next(
            e for e in entradas
            if e["accion"] == "La entidad: modifica la cantidad" and e["descripcion_linea"].startswith("Ingeniería Informática")
        )
        assert (modificacion["valor_anterior"], modificacion["valor_nuevo"], modificacion["comentario"]) == ("80", "100", "Hay demanda")
        assert modificacion["usuario_nombre"] == "admin_fac"
        assert entradas[0]["accion"] == "Generó la versión final 1" and entradas[0]["version"] == 1

    def test_los_locales_ven_lo_general_y_lo_suyo(self, api_como, admin_cum, en_aprobacion):
        entradas = api_como(admin_cum).get(f"{URL}{en_aprobacion.pk}/historial/").json()
        lineas = {e["descripcion_linea"] for e in entradas if e["descripcion_linea"]}

        assert "Envió el plan a las entidades" in [e["accion"] for e in entradas]
        assert lineas == {"Derecho — CUM Calixto García"}  # Su propio cambio, no los de la facultad.


@pytest.fixture
def publicado(en_aprobacion, vicerrector):
    return services.aprobar(en_aprobacion, vicerrector)


class TestVersionComplementaria:
    def test_parte_de_lo_aprobado_y_tiene_su_flujo(self, api_como, jefe_ci, publicado):
        datos = api_como(jefe_ci).post(f"{URL}{publicado.pk}/complementaria/", {"texto": "Ampliación"}, format="json").json()

        complementaria = PlanPlazas.objects.get(pk=datos["id"])
        assert (datos["estado"], datos["es_complementario"]) == ("borrador", True)
        assert sorted(complementaria.lineas.values_list("plazas_propuestas", flat=True)) == [30, 40, 100]

    def test_solo_una_en_curso(self, jefe_ci, publicado):
        services.crear_complementaria(publicado, jefe_ci)

        with pytest.raises(Exception, match="en curso"):
            services.crear_complementaria(publicado, jefe_ci)

    def test_al_publicarse_pasa_a_ser_la_vigente(self, jefe_ci, vicerrector, publicado, facultad, admin_fac, admin_cum, crear_usuario):
        complementaria = services.crear_complementaria(publicado, jefe_ci)
        services.actualizar_linea_ci(linea(complementaria, "II", facultad), 120, jefe_ci)
        services.enviar(complementaria, jefe_ci)
        services.cerrar_revision(complementaria, jefe_ci, forzar=True)
        services.generar_version_final(complementaria, jefe_ci)
        services.aprobar(complementaria, vicerrector)

        assert services.plan_vigente(ANIO) == complementaria
        assert OfertaCarrera.objects.get(carrera__codigo="II", facultad=facultad).plazas_plan == 120
        assert not services.admite_complementaria(publicado)
        assert services.admite_complementaria(complementaria)


class TestExportacion:
    def test_excel_del_plan_publicado(self, api_como, jefe_ci, publicado):
        response = api_como(jefe_ci).get(f"{URL}{publicado.pk}/exportar-excel/")

        hoja = load_workbook(BytesIO(response.content)).active
        filas = [[c.value for c in fila] for fila in hoja.iter_rows(min_row=6)]
        assert ["Facultad de Informática", "Facultad", "Ingeniería Informática", 100] in filas
        assert len(filas) == 3
        assert "Aprobado por" in hoja["A3"].value

    def test_solo_se_exporta_publicado(self, api_como, jefe_ci, en_aprobacion):
        response = api_como(jefe_ci).get(f"{URL}{en_aprobacion.pk}/exportar-excel/")

        assert "publicado" in errores(response)["mensaje"]

    @pytest.mark.skipif(not pdf_disponible(), reason="WeasyPrint sin librerías nativas (Pango/GTK)")
    def test_pdf_con_membrete_y_sello(self, api_como, jefe_ci, publicado):
        response = api_como(jefe_ci).get(f"{URL}{publicado.pk}/exportar-pdf/")

        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")

    @pytest.mark.skipif(pdf_disponible(), reason="Solo aplica cuando faltan las librerías de WeasyPrint")
    def test_sin_librerias_el_pdf_avisa(self, api_como, jefe_ci, publicado):
        response = api_como(jefe_ci).get(f"{URL}{publicado.pk}/exportar-pdf/")

        assert "PDF no está disponible" in errores(response)["mensaje"]

    def test_la_exportacion_se_audita(self, api_como, jefe_ci, publicado):
        from apps.auditoria.models import LogAuditoria

        api_como(jefe_ci).get(f"{URL}{publicado.pk}/exportar-excel/")

        assert LogAuditoria.objects.filter(accion="exportar plan de plazas (Excel)").exists()


def test_historial_por_cada_paso(en_aprobacion):
    assert HistorialPlan.objects.filter(plan=en_aprobacion).count() >= 8
    assert LineaPlan.objects.filter(plan=en_aprobacion).count() == 4
