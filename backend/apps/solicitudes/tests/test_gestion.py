from decimal import Decimal

import pytest
from django.core import mail
from django.utils import timezone

from apps.aspirantes import services as aspirantes
from apps.aspirantes.models import PerfilAspirante, TipoDocumento
from apps.aspirantes.tests.fixtures import cerrar, pdf
from apps.auditoria.models import LogAuditoria
from apps.core.pdf import pdf_disponible
from apps.notificaciones.models import Notificacion
from apps.solicitudes import revision, services
from apps.solicitudes.models import EventoHistorial, Observacion, Solicitud

pytestmark = pytest.mark.django_db

URL = "/api/v1/admin/solicitudes/"
OBS = "/api/v1/admin/observaciones/"


@pytest.fixture
def solicitud(aspirante, perfil_listo, convocatoria, ofertas):
    """Solicitud confirmada: Informática en la Facultad (1) y Derecho en el CUM (2)."""
    borrador = services.obtener_o_crear(aspirante)
    services.guardar_carreras(borrador, [ofertas["informatica_facultad"], ofertas["derecho_cum"]])
    return services.confirmar(borrador)


@pytest.fixture
def usuarios(crear_usuario, facultad, otra_facultad, cum):
    return {
        "sec_fac": crear_usuario("sec_fac", rol="secretario_facultad", facultad=facultad),
        "sec_cum": crear_usuario("sec_cum", rol="secretario_cum", cum_fum=cum),
        "admin_fac": crear_usuario("admin_fac", rol="admin_facultad", facultad=facultad),
        "director_cum": crear_usuario("director_cum", rol="director_cum", cum_fum=cum),
        "sec_otra": crear_usuario("sec_otra", rol="secretario_facultad", facultad=otra_facultad),
        "jefe_ci": crear_usuario("jefe_ci", rol="jefe_ci"),
    }


def tipo(nombre):
    return TipoDocumento.objects.get(nombre=nombre)


def observar(client, solicitud, **datos):
    return client.post(f"{URL}{solicitud.pk}/observar/", datos, format="json")


class TestListado:
    def test_la_entidad_ve_solo_sus_opciones(self, api_como, usuarios, solicitud):
        datos = api_como(usuarios["sec_fac"]).get(URL).json()

        assert datos["count"] == 1
        assert [c["oferta"]["entidad_nombre"] for c in datos["results"][0]["carreras"]] == ["Facultad de Informática"]
        assert datos["results"][0]["ci"] == "98042512345"

    def test_un_rol_global_ve_todas(self, api_como, usuarios, solicitud):
        carreras = api_como(usuarios["jefe_ci"]).get(URL).json()["results"][0]["carreras"]

        assert [c["prioridad"] for c in carreras] == [1, 2]

    def test_otra_entidad_no_la_ve(self, api_como, usuarios, solicitud):
        client = api_como(usuarios["sec_otra"])

        assert client.get(URL).json()["count"] == 0
        assert client.get(f"{URL}{solicitud.pk}/").status_code == 404

    def test_los_borradores_no_aparecen(self, api_como, usuarios, aspirante, convocatoria):
        services.obtener_o_crear(aspirante)

        assert api_como(usuarios["jefe_ci"]).get(URL).json()["count"] == 0

    def test_filtros(self, api_como, usuarios, solicitud, carreras, crear_aspirante, completar_perfil, ofertas):
        otra = crear_aspirante("asp_luis", ci="97010112345", email="luis@nauta.cu", first_name="Luis")
        completar_perfil(otra)
        borrador = services.obtener_o_crear(otra)
        services.guardar_carreras(borrador, [ofertas["informatica_cum"]])
        services.confirmar(borrador)
        client = api_como(usuarios["jefe_ci"])

        assert client.get(URL, {"carrera": carreras["derecho"].pk}).json()["count"] == 1
        assert client.get(URL, {"buscar": "97010112345"}).json()["results"][0]["nombre"].startswith("Luis")
        assert client.get(URL, {"indice_min": 95}).json()["count"] == 0
        assert client.get(URL, {"estado": "pendiente"}).json()["count"] == 2

    def test_filtro_por_carrera_dentro_de_su_entidad(self, api_como, usuarios, solicitud, carreras):
        # Derecho la eligió en el CUM: para la Facultad no cuenta.
        assert api_como(usuarios["sec_fac"]).get(URL, {"carrera": carreras["derecho"].pk}).json()["count"] == 0


class TestDetalle:
    def test_se_audita_y_minimiza(self, api_como, usuarios, solicitud):
        datos = api_como(usuarios["sec_cum"]).get(f"{URL}{solicitud.pk}/").json()

        assert [c["oferta"]["carrera_nombre"] for c in datos["carreras"]] == ["Derecho"]
        seleccion = next(e for e in datos["historial"] if e["tipo"] == "seleccion")
        assert [o["carrera"] for o in seleccion["despues"]] == ["Derecho"]
        assert LogAuditoria.objects.filter(accion="ver solicitud", entidad_id=str(solicitud.pk)).exists()

    def test_sin_permiso_de_descarga_no_hay_archivos(self, api_como, usuarios, solicitud):
        datos = api_como(usuarios["admin_fac"]).get(f"{URL}{solicitud.pk}/").json()

        assert all(d["documento"]["archivo"] is None for d in datos["documentos"])
        assert datos["aspirante"]["foto"] is None
        assert datos["permisos"] == {"revisar": True, "validar_indice": True, "descargar_documentos": False, "historial_documentos": False}

    def test_con_permiso_de_descarga(self, api_como, usuarios, solicitud):
        datos = api_como(usuarios["sec_fac"]).get(f"{URL}{solicitud.pk}/").json()

        assert all(d["documento"]["archivo"].startswith("/documentos/") for d in datos["documentos"])
        assert datos["aspirante"]["foto"] == f"/admin/aspirantes/{solicitud.aspirante_id}/foto/"

    def test_el_director_ve_y_descarga_pero_no_revisa(self, api_como, usuarios, solicitud):
        client = api_como(usuarios["director_cum"])
        datos = client.get(f"{URL}{solicitud.pk}/").json()

        assert datos["permisos"]["descargar_documentos"] is True
        assert datos["permisos"]["revisar"] is False
        assert client.post(f"{URL}{solicitud.pk}/validar-indice/", {"indice": 90}, format="json").status_code == 403
        assert observar(client, solicitud, objetivo="foto", texto="x").status_code == 403

    def test_elegibilidad_inicial(self, api_como, usuarios, solicitud):
        elegibilidad = api_como(usuarios["jefe_ci"]).get(f"{URL}{solicitud.pk}/").json()["elegibilidad"]

        assert elegibilidad == {"elegible": False, "motivos": ["Índice académico sin validar"]}


class TestValidarIndice:
    def url(self, solicitud):
        return f"{URL}{solicitud.pk}/validar-indice/"

    def test_valida_e_inicia_la_revision(self, api_como, usuarios, solicitud):
        response = api_como(usuarios["sec_fac"]).post(self.url(solicitud), {"indice": "91.40"}, format="json")

        assert response.status_code == 200
        assert response.json()["estado"] == "en_revision"
        assert response.json()["elegibilidad"]["elegible"] is True
        perfil = PerfilAspirante.objects.get(usuario=solicitud.aspirante)
        assert (perfil.indice_validado, perfil.indice_validado_por) == (Decimal("91.40"), usuarios["sec_fac"])
        assert EventoHistorial.objects.filter(tipo="indice_validado", usuario=usuarios["sec_fac"]).exists()

    def test_una_sola_entidad_lo_valida(self, api_como, usuarios, solicitud):
        api_como(usuarios["sec_fac"]).post(self.url(solicitud), {"indice": 91}, format="json")

        response = api_como(usuarios["sec_cum"]).post(self.url(solicitud), {"indice": 80}, format="json")

        assert response.status_code == 400
        assert "Solo la Comisión de Ingreso" in response.json()["error"]["mensaje"]

    def test_un_rol_global_puede_corregirlo(self, api_como, usuarios, solicitud):
        api_como(usuarios["sec_fac"]).post(self.url(solicitud), {"indice": 91}, format="json")

        assert api_como(usuarios["jefe_ci"]).post(self.url(solicitud), {"indice": 89}, format="json").status_code == 200
        assert PerfilAspirante.objects.get(usuario=solicitud.aspirante).indice_validado == 89

    def test_fuera_de_rango(self, api_como, usuarios, solicitud):
        assert api_como(usuarios["sec_fac"]).post(self.url(solicitud), {"indice": 101}, format="json").status_code == 400

    def test_si_el_aspirante_cambia_el_indice_se_anula(self, api_como, usuarios, solicitud, aspirante):
        api_como(usuarios["sec_fac"]).post(self.url(solicitud), {"indice": 91}, format="json")

        aspirantes.actualizar_datos(aspirante, {"indice_declarado": Decimal("93")})

        assert PerfilAspirante.objects.get(usuario=aspirante).indice_validado is None


class TestObservaciones:
    @pytest.fixture(autouse=True)
    def _on_commit(self, django_capture_on_commit_callbacks):
        self.capturar = django_capture_on_commit_callbacks

    def test_observar_la_certificacion_anula_el_indice(self, api_como, usuarios, solicitud, aspirante):
        api_como(usuarios["sec_fac"]).post(f"{URL}{solicitud.pk}/validar-indice/", {"indice": 91}, format="json")
        mail.outbox.clear()

        with self.capturar(execute=True):
            response = observar(
                api_como(usuarios["sec_fac"]), solicitud,
                objetivo="documento", tipo_documento=tipo("Certificación de estudios").pk, texto="No coincide el índice con la nota.",
            )

        assert response.status_code == 201
        datos = response.json()
        assert (datos["estado"], datos["revision_documental_pendiente"]) == ("observada", True)
        observacion = datos["observaciones"][0]
        assert (observacion["afecta_elegibilidad"], observacion["estado"], observacion["entidad_nombre"]) == (True, "abierta", "Facultad de Informática")
        assert PerfilAspirante.objects.get(usuario=aspirante).indice_validado is None
        assert Notificacion.objects.filter(usuario=aspirante, tipo="solicitud_observada").exists()
        assert "Certificación de estudios" in mail.outbox[0].body

    def test_por_defecto_el_ci_no_afecta_la_elegibilidad(self, api_como, usuarios, solicitud):
        api_como(usuarios["sec_fac"]).post(f"{URL}{solicitud.pk}/validar-indice/", {"indice": 91}, format="json")

        observar(api_como(usuarios["sec_fac"]), solicitud, objetivo="documento", tipo_documento=tipo("CI anverso").pk, texto="Ilegible")

        assert PerfilAspirante.objects.get(usuario=solicitud.aspirante).indice_validado == 91
        assert Solicitud.objects.get(pk=solicitud.pk).revision_documental_pendiente is False

    def test_la_entidad_puede_cambiar_el_valor_por_defecto(self, api_como, usuarios, solicitud):
        response = observar(api_como(usuarios["sec_fac"]), solicitud, objetivo="foto", texto="Foto de otra persona", afecta_elegibilidad=True)

        assert response.json()["observaciones"][0]["afecta_elegibilidad"] is True
        assert response.json()["revision_documental_pendiente"] is True

    def test_no_se_repite_sobre_el_mismo_elemento(self, api_como, usuarios, solicitud):
        observar(api_como(usuarios["sec_fac"]), solicitud, objetivo="foto", texto="Borrosa")

        response = observar(api_como(usuarios["sec_cum"]), solicitud, objetivo="foto", texto="Oscura")

        assert response.status_code == 400

    def test_documento_exige_tipo(self, api_como, usuarios, solicitud):
        assert observar(api_como(usuarios["sec_fac"]), solicitud, objetivo="documento", texto="x").status_code == 400

    def test_todas_las_entidades_ven_todas(self, api_como, usuarios, solicitud):
        observar(api_como(usuarios["sec_fac"]), solicitud, objetivo="foto", texto="Borrosa")

        observaciones = api_como(usuarios["sec_cum"]).get(f"{URL}{solicitud.pk}/").json()["observaciones"]

        assert [(o["entidad_nombre"], o["puede_gestionar"]) for o in observaciones] == [("Facultad de Informática", False)]

    def test_saca_la_solicitud_del_ranking(self, usuarios, solicitud):
        revision.validar_indice(usuarios["sec_fac"], solicitud, Decimal("91"))
        revision.crear_observacion(usuarios["sec_fac"], solicitud, objetivo="foto", texto="Borrosa")

        assert "Tiene observaciones sin resolver" in revision.motivos_no_elegible(Solicitud.objects.get(pk=solicitud.pk))


class TestCorreccion:
    @pytest.fixture
    def observada(self, usuarios, solicitud, convocatoria):
        """Convocatoria ya cerrada y la certificación observada."""
        revision.crear_observacion(
            usuarios["sec_fac"], solicitud, objetivo="documento", tipo_documento=tipo("Certificación de estudios"), texto="Ilegible",
        )
        cerrar(convocatoria)
        return Observacion.objects.get()

    def test_desbloqueo_selectivo_tras_el_cierre(self, api_como, aspirante, observada, ofertas):
        client = api_como(aspirante)

        corregir = client.post("/api/v1/aspirante/documentos/", {"tipo": tipo("Certificación de estudios").pk, "archivo": pdf()}, format="multipart")
        otro = client.post("/api/v1/aspirante/documentos/", {"tipo": tipo("CI anverso").pk, "archivo": pdf()}, format="multipart")
        indice = client.patch("/api/v1/aspirante/perfil/", {"indice_declarado": 99}, format="json")
        carreras = client.put(f"/api/v1/solicitudes/{observada.solicitud_id}/carreras/", {"ofertas": [ofertas["derecho_cum"].pk]}, format="json")

        assert corregir.status_code == 201
        assert corregir.json()["a_corregir"] == [{"objetivo": "documento", "tipo_documento": tipo("Certificación de estudios").pk}]
        assert (otro.status_code, indice.status_code, carreras.status_code) == (400, 400, 400)

    def test_enviar_correccion(self, api_como, aspirante, observada, usuarios):
        response = api_como(aspirante).post(f"/api/v1/solicitudes/{observada.solicitud_id}/enviar-correccion/")

        assert response.status_code == 200
        assert response.json()["estado"] == "pendiente"
        assert response.json()["observaciones"][0]["estado"] == "corregida"
        assert Notificacion.objects.filter(usuario=usuarios["sec_fac"], tipo="observacion_corregida").exists()
        # Corregida no es resuelta: sigue fuera del ranking.
        assert "Tiene observaciones sin resolver" in revision.motivos_no_elegible(Solicitud.objects.get(pk=observada.solicitud_id))

    def test_sin_observaciones_no_hay_que_corregir(self, api_como, aspirante, solicitud):
        assert api_como(aspirante).post(f"/api/v1/solicitudes/{solicitud.pk}/enviar-correccion/").status_code == 400

    def test_reabrir(self, api_como, aspirante, observada, usuarios):
        revision.enviar_correccion(observada.solicitud)

        response = api_como(usuarios["sec_fac"]).post(f"{OBS}{observada.pk}/reabrir/", {"motivo": "Sigue ilegible"}, format="json")

        assert response.status_code == 200
        assert response.json()["estado"] == "reabierta"
        assert Solicitud.objects.get(pk=observada.solicitud_id).estado == "observada"
        assert EventoHistorial.objects.filter(tipo="observacion_reabierta", despues={"motivo": "Sigue ilegible"}).exists()

    def test_solo_se_reabre_una_corregida(self, api_como, observada, usuarios):
        assert api_como(usuarios["sec_fac"]).post(f"{OBS}{observada.pk}/reabrir/", {"motivo": "x"}, format="json").status_code == 400

    def test_resolver_y_volver_a_validar(self, api_como, aspirante, observada, usuarios):
        revision.enviar_correccion(observada.solicitud)
        client = api_como(usuarios["sec_fac"])

        response = client.post(f"{OBS}{observada.pk}/resolver/")
        assert (response.status_code, response.json()["estado"]) == (200, "resuelta")
        assert Solicitud.objects.get(pk=observada.solicitud_id).estado == "en_revision"

        detalle = client.post(f"{URL}{observada.solicitud_id}/validar-indice/", {"indice": 92.5}, format="json").json()
        assert detalle["revision_documental_pendiente"] is False
        assert detalle["elegibilidad"] == {"elegible": True, "motivos": []}

    def test_otra_entidad_no_gestiona_la_ajena(self, api_como, observada, usuarios):
        client = api_como(usuarios["sec_cum"])

        assert client.post(f"{OBS}{observada.pk}/resolver/").status_code == 400
        assert client.patch(f"{OBS}{observada.pk}/", {"texto": "otro"}, format="json").status_code == 400

    def test_un_rol_global_gestiona_cualquiera(self, api_como, observada, usuarios):
        response = api_como(usuarios["jefe_ci"]).patch(f"{OBS}{observada.pk}/", {"texto": "Falta el cuño"}, format="json")

        assert (response.status_code, response.json()["texto"]) == (200, "Falta el cuño")
        assert EventoHistorial.objects.filter(tipo="observacion_editada").exists()

    def test_historial_del_documento(self, aspirante, observada):
        aspirantes.subir_documento(aspirante, tipo("Certificación de estudios"), pdf("nueva.pdf"), "PDF")

        evento = EventoHistorial.objects.filter(tipo="documento").first()
        assert evento.descripcion == "Sustituyó Certificación de estudios"
        assert (evento.antes["version"], evento.despues) == (1, {"version": 2, "archivo": "nueva.pdf"})
        assert evento.solicitud_id == observada.solicitud_id


class TestEstadisticas:
    def test_termometro_por_oferta_y_ambito(self, api_como, usuarios, solicitud, ofertas):
        datos = api_como(usuarios["jefe_ci"]).get(f"{URL}estadisticas/").json()

        fila = next(f for f in datos["termometro"] if f["oferta"] == ofertas["derecho_cum"].pk)
        assert (fila["plazas"], fila["solicitudes"], fila["primera_opcion"]) == (5, 1, 0)
        assert datos["totales"]["solicitudes"] == 1
        assert datos["totales"]["por_estado"]["pendiente"] == 1
        assert datos["evolucion"] == [{"dia": str(timezone.localdate()), "total": 1}]

        local = api_como(usuarios["sec_fac"]).get(f"{URL}estadisticas/").json()
        assert {f["entidad"] for f in local["termometro"]} == {"Facultad de Informática"}


class TestExportar:
    def test_excel_auditado_con_sus_opciones(self, api_como, usuarios, solicitud):
        from io import BytesIO

        from openpyxl import load_workbook

        response = api_como(usuarios["sec_cum"]).get(f"{URL}exportar-excel/")

        assert response.status_code == 200
        hoja = load_workbook(BytesIO(response.content)).active
        textos = [str(c.value) for fila in hoja.iter_rows() for c in fila if c.value]
        assert any("Derecho" in t for t in textos)
        assert not any("Informática" in t for t in textos)
        assert LogAuditoria.objects.filter(accion="exportar solicitudes (Excel)").exists()

    def test_sin_permiso_de_exportar(self, api_como, crear_usuario, facultad, solicitud):
        from apps.nomencladores.models import Area

        jefe = crear_usuario("jefe_area", rol="jefe_area_facultad", facultad=facultad, area=Area.objects.create(nombre="A", facultad=facultad))

        assert api_como(jefe).get(f"{URL}exportar-excel/").status_code == 403

    @pytest.mark.skipif(not pdf_disponible(), reason="WeasyPrint sin GTK en este equipo")
    def test_pdf(self, api_como, usuarios, solicitud):
        response = api_como(usuarios["jefe_ci"]).get(f"{URL}exportar-pdf/")

        assert response.status_code == 200
        assert response.content.startswith(b"%PDF")
